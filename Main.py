
import os
import re

from pypdf import PdfReader
from docx import Document


# ==================================================
# SKILL DICTIONARY
# Add new skills here whenever needed.
# ==================================================

SKILL_ALIASES = {
    "Python": [r"\bpython\b"],
    "JavaScript": [r"\bjavascript\b"],
    "TypeScript": [r"\btypescript\b"],
    "SQL": [r"\bsql\b"],
    "MySQL": [r"\bmysql\b"],
    "PostgreSQL": [r"\bpostgresql\b"],
    "MongoDB": [r"\bmongodb\b"],
    "SQLite": [r"\bsqlite\b"],

    # Data Science and AI
    "NumPy": [r"\bnumpy\b"],
    "Pandas": [r"\bpandas\b"],
    "SciPy": [r"\bscipy\b"],
    "Matplotlib": [r"\bmatplotlib\b"],
    "Seaborn": [r"\bseaborn\b"],
    "PySpark": [r"\bpyspark\b"],
    "Machine Learning": [r"\bmachine learning\b"],
    "Deep Learning": [r"\bdeep learning\b"],
    "Artificial Intelligence": [
        r"\bartificial intelligence\b",
        r"\bAI\b",
    ],
    "Generative AI": [r"\bgenerative AI\b"],
    "Agentic AI": [r"\bagentic AI\b"],
    "Natural Language Processing": [
        r"\bnatural language processing\b",
        r"\bNLP\b",
    ],
    "Computer Vision": [r"\bcomputer vision\b"],
    "OpenCV": [r"\bopencv\b"],
    "TensorFlow": [r"\btensorflow\b"],
    "PyTorch": [r"\bpytorch\b"],
    "Scikit-learn": [
        r"\bscikit.learn\b",
        r"\bsklearn\b",
    ],
    "XGBoost": [r"\bxgboost\b"],
    "Neural Networks": [r"\bneural networks?\b"],
    "CNN": [r"\bCNN\b"],
    "ANN": [r"\bANN\b"],
    "ARIMA": [r"\barima\b"],
    "Statistics": [r"\bstatistics\b"],
    "Data Analysis": [
        r"\bdata analysis\b",
        r"\bdata analytics\b",
    ],
    "Data Visualization": [r"\bdata visualization\b"],
    "Exploratory Data Analysis": [
        r"\bexploratory data analysis\b",
        r"\bEDA\b",
    ],
    "Data Cleaning": [r"\bdata cleaning\b"],
    "Feature Engineering": [r"\bfeature engineering\b"],
    "Regression": [r"\bregression\b"],
    "Classification": [r"\bclassification\b"],
    "Clustering": [r"\bclustering\b"],
    "Random Forest": [r"\brandom forest\b"],
    "Decision Trees": [r"\bdecision trees?\b"],
    "Time-Series Forecasting": [
        r"\btime.series forecasting\b",
    ],
    "Probability": [r"\bprobability\b"],
    "Hypothesis Testing": [r"\bhypothesis testing\b"],

    # Web Development
    "HTML": [r"\bhtml5?\b"],
    "CSS": [r"\bcss3?\b"],
    "React": [r"\breact(?:\.js)?\b"],
    "Angular": [r"\bangular\b"],
    "Vue.js": [r"\bvue(?:\.js)?\b"],
    "Flask": [r"\bflask\b"],
    "FastAPI": [r"\bfastapi\b"],
    "Django": [r"\bdjango\b"],
    "Node.js": [r"\bnode\.?js\b"],
    "Express.js": [r"\bexpress\.?js\b"],
    "RESTful APIs": [
        r"\bRESTful APIs?\b",
        r"\bREST APIs?\b",
    ],
    "Bootstrap": [r"\bbootstrap\b"],
    "Tailwind CSS": [r"\btailwind\s*css\b"],
    "Responsive Design": [r"\bresponsive design\b"],
    "Form Validation": [r"\bform validation\b"],

    # Programming concepts
    "Data Structures and Algorithms": [
        r"\bdata structures?\s*(?:and|&)\s*algorithms?\b",
        r"\bDSA\b",
    ],
    "Object-Oriented Programming": [
        r"\bobject.oriented programming\b",
        r"\bOOP\b",
    ],
    "Problem Solving": [r"\bproblem.solving\b"],
    "File Handling": [r"\bfile handling\b"],
    "REST API Integration": [
        r"\bREST API integration\b",
    ],

    # Tools and Cloud
    "Git": [r"\bgit\b"],
    "GitHub": [r"\bgithub\b"],
    "Docker": [r"\bdocker\b"],
    "Kubernetes": [r"\bkubernetes\b"],
    "Linux": [r"\blinux\b"],
    "AWS": [r"\bAWS\b"],
    "Microsoft Azure": [r"\bmicrosoft azure\b", r"\bazure\b"],
    "Google Cloud": [r"\bgoogle cloud\b"],
    "Jupyter Notebook": [r"\bjupyter notebooks?\b"],
    "Google Colab": [r"\bgoogle colab\b"],
    "VS Code": [r"\bVS\s*Code\b"],
    "Raspberry Pi": [r"\braspberry pi\b"],
    "Arduino": [r"\barduino\b"],
    "IoT": [r"\bIoT\b"],

    # Business and other technologies
    "Razorpay": [r"\brazorpay\b"],
    "Pytest": [r"\bpytest\b"],
    "Selenium": [r"\bselenium\b"],
    "Power BI": [r"\bpower\s*BI\b"],
    "Tableau": [r"\btableau\b"],
    "Excel": [r"\bexcel\b"],
}


# ==================================================
# READ PDF AND DOCX RESUMES
# ==================================================

def read_resume(file_path):
    """
    Read text from a PDF or DOCX resume.
    Accepts a file path and returns extracted text.
    """

    if not file_path or not os.path.isfile(file_path):
        raise FileNotFoundError(
            f"Resume file not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()
    extracted_text = []

    if extension == ".pdf":
        reader = PdfReader(file_path)

        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_text.append(page_text)

    elif extension == ".docx":
        document = Document(file_path)

        # Read paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                extracted_text.append(paragraph.text)

        # Read tables, which are common in resumes
        for table in document.tables:
            for row in table.rows:
                row_text = " ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )
                if row_text:
                    extracted_text.append(row_text)

    else:
        raise ValueError(
            "Unsupported file format. Please upload PDF or DOCX."
        )

    resume_text = "\n".join(extracted_text).strip()

    if not resume_text:
        raise ValueError(
            "No readable text found. This may be a scanned PDF. "
            "Use OCR to read image-based resumes."
        )

    return resume_text


# ==================================================
# EXTRACT SKILLS FROM THE ACTUAL RESUME TEXT
# ==================================================

def extract_skills(resume_text):
    """
    Detect supported skills mentioned in the supplied resume text.
    Does not add a skill unless its pattern is found in the text.
    """

    if not isinstance(resume_text, str) or not resume_text.strip():
        return []

    detected_skills = []

    for skill, patterns in SKILL_ALIASES.items():
        for pattern in patterns:
            if re.search(
                pattern,
                resume_text,
                flags=re.IGNORECASE,
            ):
                detected_skills.append(skill)
                break

    return detected_skills


# ==================================================
# DISPLAY SKILL NAMES
# ==================================================

def display_skill(skill):
    """Return a consistent display name for a detected skill."""

    if not skill:
        return ""

    return skill


# ==================================================
# OPTIONAL: EXTRACT A TEXT SECTION BY HEADING
# ==================================================

def extract_section(resume_text, start_heading, end_headings=None):
    """
    Extract a section between headings.
    Useful for Projects, Experience, Education, etc.
    """

    if not resume_text:
        return ""

    if end_headings is None:
        end_headings = [
            "experience",
            "work experience",
            "education",
            "skills",
            "technical skills",
            "certifications",
            "achievements",
            "projects",
            "summary",
            "professional summary",
        ]

    lines = resume_text.splitlines()
    start_pattern = re.compile(
        rf"^\s*{re.escape(start_heading)}\s*:?\s*$",
        re.IGNORECASE,
    )

    inside_section = False
    collected_lines = []

    for line in lines:
        cleaned = line.strip()

        if not inside_section:
            if start_pattern.match(cleaned):
                inside_section = True
            continue

        if any(
            re.match(
                rf"^\s*{re.escape(heading)}\s*:?\s*$",
                cleaned,
                re.IGNORECASE,
            )
            for heading in end_headings
        ):
            break

        if cleaned:
            collected_lines.append(cleaned)

    return "\n".join(collected_lines)


# ==================================================
# LOCAL TEST
# Run: python Main.py "path/to/resume.pdf"
# ==================================================

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print(
            'Usage: python Main.py "C:\\path\\to\\resume.pdf"'
        )
    else:
        path = sys.argv[1].strip('"')

        try:
            text = read_resume(path)
            skills = extract_skills(text)

            print("\nResume:", os.path.basename(path))
            print("Skills detected:", len(skills))

            for skill in skills:
                print("-", display_skill(skill))

        except Exception as error:
            print("Error:", error)
